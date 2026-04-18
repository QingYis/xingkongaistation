import asyncio
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches
import os

async def export_slides():
    html_path = f"file://{os.path.abspath('index.html')}"
    screenshots = []
    
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        # Create a 16:9 viewport
        page = await browser.new_page(viewport={"width": 1920, "height": 1080})
        await page.goto(html_path)
        
        # Wait for fonts and animations to settle
        await page.wait_for_timeout(2000)
        
        # Hide the nav dots so they don't appear in the PPT
        await page.evaluate("document.querySelector('.nav-dots').style.display = 'none';")
        
        # Take screenshots of each section
        slide_ids = ["slide-hero", "slide-models", "slide-launch", "slide-steps", "slide-support"]
        for i, slide_id in enumerate(slide_ids):
            element = page.locator(f"#{slide_id}")
            # Scroll to the element to ensure any scroll-triggered animations run
            await element.scroll_into_view_if_needed()
            await page.wait_for_timeout(2500) # wait longer for all fade-in animations to finish
            
            img_path = f"slide_{i+1}.png"
            await element.screenshot(path=img_path)
            screenshots.append(img_path)
            print(f"Captured {slide_id}")
            
        await browser.close()
        
    # Create PPTX
    prs = Presentation()
    # Set slide dimensions to 16:9
    prs.slide_width = Inches(16)
    prs.slide_height = Inches(9)
    
    blank_slide_layout = prs.slide_layouts[6]
    
    # Insert the user's uploaded image as the second slide
    user_image_path = "/Users/xingkong/Desktop/AIcodex/xingkongaistation/截屏2026-04-18 20.14.41.png"
    if os.path.exists(user_image_path):
        screenshots.insert(1, user_image_path)
    
    for img_path in screenshots:
        slide = prs.slides.add_slide(blank_slide_layout)
        slide.shapes.add_picture(img_path, 0, 0, width=Inches(16), height=Inches(9))
        print(f"Added {img_path} to PPT")
        
    prs.save("../星空AI宣传PPT_视觉重构版.pptx")
    print("Saved as 星空AI宣传PPT_视觉重构版.pptx")
    
    # Cleanup images
    for img_path in screenshots:
        if img_path != user_image_path and os.path.exists(img_path):
            os.remove(img_path)

if __name__ == "__main__":
    asyncio.run(export_slides())
